"""Receipt Scanner: turns photos of shop bills into a checked expense
spreadsheet.

Claude reads each photo and copies out the details. Then plain Python code
(not the model) checks the arithmetic, spots duplicate bills and writes
expenses.xlsx. Run it with:  python agent.py
"""

import asyncio
import json
from pathlib import Path

from claude_agent_sdk import ClaudeAgentOptions, ResultMessage, query
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from watch import show

HERE = Path(__file__).parent

SCHEMA = {
    "type": "object",
    "properties": {
        "receipts": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "file": {
                        "type": "string",
                        "description": "File name, like receipt-01.jpg",
                    },
                    "shop": {"type": "string"},
                    "bill_no": {"type": "string"},
                    "date": {
                        "type": "string",
                        "description": "DD/MM/YYYY, as printed",
                    },
                    "gstin": {
                        "type": "string",
                        "description": "Empty if none is printed",
                    },
                    "items": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "name": {"type": "string"},
                                "qty": {"type": "number"},
                                "rate": {"type": "number"},
                                "amount": {"type": "number"},
                            },
                            "required": ["name", "qty", "rate", "amount"],
                        },
                    },
                    "subtotal": {"type": "number"},
                    "gst_amount": {
                        "type": "number",
                        "description": "CGST + SGST, 0 if none",
                    },
                    "total": {
                        "type": "number",
                        "description": "The TOTAL exactly as printed",
                    },
                    "category": {
                        "enum": [
                            "groceries",
                            "food",
                            "hardware",
                            "medical",
                            "business supplies",
                            "stationery",
                            "fuel",
                            "other",
                        ]
                    },
                },
                "required": [
                    "file",
                    "shop",
                    "bill_no",
                    "date",
                    "gstin",
                    "items",
                    "subtotal",
                    "gst_amount",
                    "total",
                    "category",
                ],
            },
        }
    },
    "required": ["receipts"],
}

options = ClaudeAgentOptions(
    model="claude-sonnet-5-5",
    cwd=str(HERE),
    system_prompt=(
        "You copy details from photos of Indian shop receipts. Look at "
        "every image in the receipts folder. Copy numbers exactly as "
        "printed, even if they look wrong: don't correct or recalculate "
        "anything."
    ),
    tools=["Read", "Glob"],
    allowed_tools=["Read", "Glob"],
    output_format={"type": "json_schema", "schema": SCHEMA},
    max_turns=30,
    max_budget_usd=1.00,
)


def check(receipts):
    """Plain-code checks. Each receipt gets a list of problems (empty means
    OK)."""
    seen = {}
    for r in receipts:
        r["file"] = Path(r["file"]).name  # "receipts/x.jpg" -> "x.jpg"
        problems = []
        items_sum = round(sum(i["amount"] for i in r["items"]), 2)
        if abs(items_sum - r["subtotal"]) > 0.5:
            problems.append(
                f"items add up to {items_sum:.2f}, not {r['subtotal']:.2f}"
            )
        expected = round(r["subtotal"] + r["gst_amount"], 2)
        if abs(expected - r["total"]) > 0.5:
            problems.append(
                f"total should be {expected:.2f} but bill says "
                f"{r['total']:.2f}"
            )
        key = (r["shop"].lower(), r["bill_no"].lower(), r["date"])
        if key in seen:
            problems.append(f"duplicate of {seen[key]}")
        seen.setdefault(key, r["file"])
        r["problems"] = problems
    return receipts


def write_excel(receipts, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses"
    ws.append(
        [
            "File",
            "Date",
            "Shop",
            "Bill no",
            "GSTIN",
            "Category",
            "Subtotal",
            "GST",
            "Total",
            "Check",
        ]
    )
    for c in ws[1]:
        c.font = Font(bold=True)
    red = PatternFill("solid", fgColor="FFD7D7")
    for r in sorted(receipts, key=lambda r: r["file"]):
        ws.append(
            [
                r["file"],
                r["date"],
                r["shop"],
                r["bill_no"],
                r["gstin"],
                r["category"],
                r["subtotal"],
                r["gst_amount"],
                r["total"],
                "; ".join(r["problems"]) or "OK",
            ]
        )
        if r["problems"]:
            for c in ws[ws.max_row]:
                c.fill = red
    good = [r for r in receipts if not r["problems"]]
    ws.append([])
    ws.append(
        [
            "",
            "",
            "Total of checked bills",
            "",
            "",
            "",
            round(sum(r["subtotal"] for r in good), 2),
            round(sum(r["gst_amount"] for r in good), 2),
            round(sum(r["total"] for r in good), 2),
            f"{len(good)} of {len(receipts)} OK",
        ]
    )
    for col, width in zip(
        "ABCDEFGHIJ", [16, 12, 30, 13, 18, 17, 11, 9, 11, 48]
    ):
        ws.column_dimensions[col].width = width
    wb.save(path)


async def main():
    async for message in query(
        prompt="Read every receipt photo in the receipts folder.",
        options=options,
    ):
        show(message)
        if isinstance(message, ResultMessage) and message.structured_output:
            receipts = check(message.structured_output["receipts"])
            (HERE / "receipts.json").write_text(
                json.dumps(receipts, indent=2)
            )
            write_excel(receipts, HERE / "expenses.xlsx")
            for r in receipts:
                print(
                    f"{r['file']}: {r['shop']}, Rs. {r['total']:,.2f} -> "
                    f"{'; '.join(r['problems']) or 'OK'}"
                )


if __name__ == "__main__":
    asyncio.run(main())
